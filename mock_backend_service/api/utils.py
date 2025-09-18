import io
import json
from dataclasses import dataclass
from typing import Any, Dict, List, Optional, Tuple
import requests
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet


@dataclass
class EndpointSpec:
    """Represents an endpoint discovered in swagger/openapi."""
    method: str
    path: str
    consumes: Optional[List[str]] = None
    produces: Optional[List[str]] = None
    base_url: Optional[str] = None


# PUBLIC_INTERFACE
def parse_swagger_json(swagger_bytes: bytes) -> Dict[str, EndpointSpec]:
    """
    PUBLIC_INTERFACE
    Parse a swagger.json (OpenAPI 2.0) file and return a mapping of operationId (or method+path)
    to EndpointSpec objects. This implementation handles basic swagger 2.0 structures.
    """
    text = swagger_bytes.decode("utf-8").strip()
    data = json.loads(text)
    host = data.get("host", "")
    schemes = data.get("schemes", ["http"])
    base_path = data.get("basePath", "")
    base_url = None
    if host:
        scheme = schemes[0] if schemes else "http"
        base_url = f"{scheme}://{host}{base_path}"
    elif base_path:
        # If only basePath is provided, use relative base
        base_url = base_path
    else:
        base_url = ""

    out: Dict[str, EndpointSpec] = {}
    paths = data.get("paths", {})
    global_consumes = data.get("consumes")
    global_produces = data.get("produces")

    for path, methods in paths.items():
        for method, op in methods.items():
            if method.lower() not in ("get", "post", "put", "patch", "delete", "options", "head"):
                continue
            op_id = op.get("operationId") or f"{method.lower()}_{path}".replace("/", "_")
            consumes = op.get("consumes", global_consumes)
            produces = op.get("produces", global_produces)
            out[op_id] = EndpointSpec(
                method=method.upper(),
                path=path,
                consumes=consumes,
                produces=produces,
                base_url=base_url,
            )
    return out


def _normalize_url(base_url: str, path: str) -> str:
    if base_url.endswith("/"):
        base_url = base_url[:-1]
    if not path.startswith("/"):
        path = "/" + path
    return f"{base_url}{path}"


def _parse_row_dict(row_values: Dict[str, Any]) -> Tuple[str, str, Dict[str, Any]]:
    """
    Expect at least columns: method, path or operationId, payload (JSON as string).
    If base_url is provided in a column, it overrides swagger host/basePath.
    """
    operation_id = str(row_values.get("operationid") or row_values.get("operation_id") or "").strip()
    method = str(row_values.get("method") or "").strip().upper()
    path = str(row_values.get("path") or "").strip()
    payload_text = row_values.get("payload") or row_values.get("body") or "{}"
    try:
        payload = json.loads(payload_text) if isinstance(payload_text, str) else (payload_text or {})
    except Exception:
        payload = {}
    base_url_override = str(row_values.get("base_url") or "").strip()
    # Expected status may be missing or invalid; mark as insufficient later if invalid
    expected_status_raw = row_values.get("expected_status")
    try:
        expected_status = int(expected_status_raw) if expected_status_raw not in (None, "") else 200
    except Exception:
        expected_status = None  # invalid, will trigger Data Insufficient
    return operation_id, base_url_override, {
        "method": method, "path": path, "payload": payload, "expected_status": expected_status
    }


# PUBLIC_INTERFACE
def execute_testcases(swagger_bytes: bytes, testcase_bytes: bytes) -> Tuple[bytes, int, int, int]:
    """
    PUBLIC_INTERFACE
    Execute test cases from an Excel against endpoints defined in swagger.

    IMPORTANT: Only the worksheet named 'Standard Template' is processed. Other sheets
    are left untouched and unmodified.

    Required columns on 'Standard Template' (case-insensitive headers):
      - method (e.g., GET, POST) or operationId (to match swagger operationId)
      - path (e.g., /health/) when method is specified (ignored if operationId uniquely maps)
      - payload (JSON string) for body where applicable
      - expected_status (optional, default 200)
      - base_url (optional to override swagger base)

    Adds/Updates a 'Status' column on the 'Standard Template' sheet with:
      - 'Pass' if actual status matches expected
      - 'Fail' if actual status doesn't match expected or request errors with valid expected status
      - 'Data Insufficient' if required fields are missing/invalid (e.g., no method/opId/path, invalid expected_status)

    Returns tuple of (result_excel_bytes, total, passed, failed)
    """
    specs = parse_swagger_json(swagger_bytes)

    wb: Workbook = load_workbook(io.BytesIO(testcase_bytes))
    # Only work with 'Standard Template' and do not touch other sheets
    if "Standard Template" not in wb.sheetnames:
        # If the required sheet is missing, return the original file with zero counts
        out_stream = io.BytesIO()
        wb.save(out_stream)
        out_stream.seek(0)
        return out_stream.read(), 0, 0, 0

    ws: Worksheet = wb["Standard Template"]

    # Build header map (case-insensitive)
    headers: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header = str(ws.cell(row=1, column=col).value or "").strip()
        if header:
            headers[header.lower()] = col

    # Ensure Status column exists on Standard Template only
    status_col = headers.get("status")
    if not status_col:
        status_col = ws.max_column + 1
        ws.cell(row=1, column=status_col, value="Status")
        headers["status"] = status_col

    total = 0
    passed = 0
    failed = 0

    # Determine last row; iterate through all rows but skip completely empty ones
    last_row = ws.max_row

    for row in range(2, last_row + 1):
        # Consider a row empty if all non-status header columns are empty
        is_empty = True
        row_values: Dict[str, Any] = {}
        for h, c in headers.items():
            if h == "status":
                continue
            val = ws.cell(row=row, column=c).value
            row_values[h] = val
            if val not in (None, ""):
                is_empty = False

        if is_empty:
            # Do not write anything for completely empty rows
            continue

        total += 1

        operation_id, base_url_override, parsed = _parse_row_dict(row_values)
        method = parsed["method"]
        path = parsed["path"]
        payload = parsed["payload"]
        expected_status = parsed["expected_status"]

        # Validate row data sufficiency
        data_insufficient = False
        if expected_status is None:
            data_insufficient = True
        if not operation_id and not method:
            data_insufficient = True
        if not operation_id and method and not path:
            data_insufficient = True

        # Resolve URL and HTTP method
        url: Optional[str] = None
        if not data_insufficient:
            if operation_id:
                spec = specs.get(operation_id)
                if not spec:
                    # Unknown operationId => insufficient mapping info
                    data_insufficient = True
                else:
                    url = _normalize_url(base_url_override or spec.base_url or "", spec.path)
                    method = spec.method
            else:
                # fall back to method+path
                found_spec = None
                for sp in specs.values():
                    if sp.method == method and sp.path == path:
                        found_spec = sp
                        break
                if found_spec:
                    url = _normalize_url(base_url_override or found_spec.base_url or "", found_spec.path)
                else:
                    # if base_url provided, allow direct URL build; otherwise insufficient target
                    if base_url_override:
                        url = _normalize_url(base_url_override, path)
                    else:
                        data_insufficient = True

        if data_insufficient:
            ws.cell(row=row, column=status_col, value="Data Insufficient")
            continue

        status_value = "Fail"
        try:
            # Send request
            if method == "GET":
                resp = requests.get(url, params=payload if isinstance(payload, dict) else None, timeout=15)
            elif method in ("POST", "PUT", "PATCH", "DELETE"):
                resp = requests.request(method, url, json=payload, timeout=15)
            else:
                resp = requests.request(method, url, timeout=15)

            if resp.status_code == expected_status:
                status_value = "Pass"
                passed += 1
            else:
                failed += 1
        except Exception:
            # Only mark as Fail if we had sufficient data and attempted the request
            failed += 1
            status_value = "Fail"

        ws.cell(row=row, column=status_col, value=status_value)

    # Save workbook back to bytes (only 'Standard Template' modified; others untouched)
    out_stream = io.BytesIO()
    wb.save(out_stream)
    out_stream.seek(0)
    return out_stream.read(), total, passed, failed
