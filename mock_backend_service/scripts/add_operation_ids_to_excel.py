#!/usr/bin/env python3
"""
PUBLIC_INTERFACE
Utility Script: Populate 'operationId' column in an Excel file from an OpenAPI/Swagger spec.

This script:
- Reads operationIds from a swagger file (JSON or YAML; .txt allowed with JSON content).
- Opens an Excel workbook and targets the 'Standard Template' sheet.
- Adds a new column named 'operationId' if not present.
- Populates that column with all operationIds found in the swagger file (one per row).
  - If more operationIds than available rows, it fills up to the last existing row.
  - If more rows than operationIds, it leaves remaining cells blank.
- Saves the modified workbook as a new file.

Default paths (can be overridden via CLI args):
  - swagger: mock_backend_service/inputs/swagger.txt
  - input  : mock_backend_service/inputs/testcases.xlsx
  - output : mock_backend_service/inputs/testcases_with_operationId.xlsx
  - sheet  : "Standard Template"
  - column : "operationId"

Example:
  python mock_backend_service/scripts/add_operation_ids_to_excel.py
  python mock_backend_service/scripts/add_operation_ids_to_excel.py \
      --swagger mock_backend_service/inputs/swagger.txt \
      --input   mock_backend_service/inputs/testcases.xlsx \
      --output  mock_backend_service/inputs/testcases_with_operationId.xlsx \
      --sheet "Standard Template" \
      --column "operationId"
"""

import argparse
import json
import os
from typing import Any, Dict, Iterable, List, Tuple

try:
    import yaml  # PyYAML for YAML specs
except Exception:
    yaml = None

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


# ----------------------------
# Parsing OpenAPI/Swagger spec
# ----------------------------

def _load_text_bytes(path: str) -> bytes:
    with open(path, "rb") as f:
        return f.read()


def _parse_spec(text: str) -> Dict[str, Any]:
    """
    Try JSON first; if that fails and yaml is available, try YAML.
    """
    try:
        return json.loads(text)
    except Exception:
        if yaml is None:
            raise RuntimeError(
                "Failed to parse as JSON and PyYAML not available. "
                "Install PyYAML or provide valid JSON."
            )
        try:
            return yaml.safe_load(text)
        except Exception as e:
            raise RuntimeError(f"Failed to parse spec as JSON or YAML: {e}") from e


def _iter_operations_from_paths(paths: Dict[str, Any]) -> Iterable[Tuple[str, str, Dict[str, Any]]]:
    """
    Iterate over HTTP operations from an OpenAPI/Swagger 'paths' object.
    Yields (method_upper, path, operation_object_dict).
    """
    if not isinstance(paths, dict):
        return
    http_methods = {"get", "post", "put", "patch", "delete", "head", "options", "trace"}
    for path, path_item in paths.items():
        if not isinstance(path_item, dict):
            continue
        for method, op_obj in path_item.items():
            if method.lower() in http_methods and isinstance(op_obj, dict):
                yield method.upper(), path, op_obj


# PUBLIC_INTERFACE
def extract_operation_ids_from_spec(spec: Dict[str, Any]) -> List[str]:
    """
    PUBLIC_INTERFACE
    Extract a list of operationIds from an OpenAPI/Swagger spec.
    If an operation lacks an operationId, a synthesized one is produced.

    Returns:
        List of operationId strings (unique order as encountered).
    """
    ids: List[str] = []
    seen = set()

    paths = spec.get("paths")
    if not isinstance(paths, dict):
        return ids

    for method, path, op in _iter_operations_from_paths(paths):
        op_id = op.get("operationId")
        if not op_id:
            # Synthesize a readable id when missing
            op_id = f"{method.lower()}_{path}".replace("/", "_").replace("{", "").replace("}", "")
        if op_id not in seen:
            ids.append(op_id)
            seen.add(op_id)
    return ids


# ----------------------------
# Excel manipulation helpers
# ----------------------------

def _get_or_create_column(ws: Worksheet, header_name: str) -> int:
    """
    Return the column index for a case-insensitive header 'header_name'.
    If not found, create a new header in the next empty column and return its index.
    """
    header_lower = header_name.strip().lower()
    # Build mapping of existing headers
    for col in range(1, ws.max_column + 1):
        cell_val = ws.cell(row=1, column=col).value
        if isinstance(cell_val, str) and cell_val.strip().lower() == header_lower:
            return col
    # Not found, create at next column
    new_col = ws.max_column + 1 if ws.max_column else 1
    ws.cell(row=1, column=new_col, value=header_name)
    return new_col


# PUBLIC_INTERFACE
def populate_operation_ids_in_excel(
    swagger_path: str,
    input_excel_path: str,
    output_excel_path: str,
    sheet_name: str = "Standard Template",
    column_name: str = "operationId",
) -> None:
    """
    PUBLIC_INTERFACE
    Read operationIds from swagger_path and populate a new/existing 'operationId' column
    in the specified sheet of the input Excel, then save to output_excel_path.

    Behavior:
    - Only the 'Standard Template' sheet is modified; other sheets are untouched.
    - If the sheet does not exist, the workbook is saved unchanged to the output path.

    Args:
        swagger_path: Path to OpenAPI/Swagger file (JSON or YAML; .txt allowed).
        input_excel_path: Path to the input Excel file (.xlsx).
        output_excel_path: Path where the modified Excel will be saved.
        sheet_name: Sheet to modify (default "Standard Template").
        column_name: Column header to add/fill (default "operationId").
    """
    if not os.path.isfile(swagger_path):
        raise FileNotFoundError(f"Swagger file not found: {swagger_path}")
    if not os.path.isfile(input_excel_path):
        raise FileNotFoundError(f"Input Excel not found: {input_excel_path}")

    # Load and parse spec
    spec_bytes = _load_text_bytes(swagger_path)
    text = spec_bytes.decode("utf-8", errors="replace")
    spec = _parse_spec(text)
    op_ids = extract_operation_ids_from_spec(spec)

    # Load workbook
    wb = load_workbook(input_excel_path)

    # If sheet missing, save workbook unchanged to output and return
    if sheet_name not in wb.sheetnames:
        wb.save(output_excel_path)
        return

    ws = wb[sheet_name]

    # Find or create the target column
    target_col = _get_or_create_column(ws, column_name)

    # Determine how many data rows exist: use ws.max_row, but skip header
    last_row = ws.max_row if ws.max_row and ws.max_row >= 2 else 1

    # Fill operationIds row-wise starting from row 2
    idx = 0
    for row in range(2, last_row + 1):
        value = op_ids[idx] if idx < len(op_ids) else None
        ws.cell(row=row, column=target_col, value=value)
        if idx < len(op_ids):
            idx += 1

    # Save as new file
    wb.save(output_excel_path)


def _build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Populate 'operationId' column in an Excel sheet from an OpenAPI/Swagger spec."
    )
    parser.add_argument(
        "--swagger",
        default="mock_backend_service/inputs/swagger.txt",
        help="Path to the swagger/openapi file (JSON or YAML). Default: mock_backend_service/inputs/swagger.txt",
    )
    parser.add_argument(
        "--input",
        dest="input_excel",
        default="mock_backend_service/inputs/testcases.xlsx",
        help="Path to input Excel (.xlsx). Default: mock_backend_service/inputs/testcases.xlsx",
    )
    parser.add_argument(
        "--output",
        dest="output_excel",
        default="mock_backend_service/inputs/testcases_with_operationId.xlsx",
        help="Path to output Excel (.xlsx). Default: mock_backend_service/inputs/testcases_with_operationId.xlsx",
    )
    parser.add_argument(
        "--sheet",
        default="Standard Template",
        help="Worksheet name to modify. Default: 'Standard Template'",
    )
    parser.add_argument(
        "--column",
        default="operationId",
        help="Column header to add/fill. Default: 'operationId'",
    )
    return parser


def main() -> int:
    parser = _build_arg_parser()
    args = parser.parse_args()

    try:
        populate_operation_ids_in_excel(
            swagger_path=args.swagger,
            input_excel_path=args.input_excel,
            output_excel_path=args.output_excel,
            sheet_name=args.sheet,
            column_name=args.column,
        )
        print(
            f"Success: Wrote operationIds to '{args.column}' in sheet '{args.sheet}' "
            f"and saved to '{args.output_excel}'."
        )
        return 0
    except Exception as exc:
        print(f"Error: {exc}")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
